"""convert_ws2: parsers + the annotator reader, against a fixture built from the
real NWFA template (invented answers — never shipped, only used in tmp_path)."""

import datetime as dt
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import convert_ws2 as c  # noqa: E402

TEMPLATE = c.META / "NWFA_annotation_TEMPLATE_v0.1.xlsx"


def test_parse_timecode_variants():
    assert c.parse_timecode("00:01:26") == 86
    assert c.parse_timecode("04:24") == 264
    assert c.parse_timecode(dt.time(0, 7, 42)) == 462
    assert c.parse_timecode(dt.timedelta(minutes=3)) == 180
    assert c.parse_timecode(None) is None and c.parse_timecode("") is None


def test_inner_segments_mixed_formats():
    assert c.parse_inner_segments("00:00: A; 04:24: B") == [(0, "A"), (264, "B")]
    assert c.parse_inner_segments("00:00:00: A; 00:22:24: B.") == [(0, "A"), (1344, "B")]


def test_ref_of():
    assert c.ref_of("NWFA 57 [GOVERNMENT PROPAGANDA FILMS]") == "NWFA 57"
    assert c.ref_of("N_6975915_01of01_Online Shopping") == "N_6975915"


def make_annotator(path: Path, name: str, split_at: int):
    wb = openpyxl.load_workbook(TEMPLATE)
    wb["Videos"]["C1"] = name
    ws = wb["Programmes"]
    # first data row = row 4 (header on row 3); template ref N_6966411 programmes 1..6
    rows = [(4, "00:00:10", f"00:{split_at:02d}:00", "Interview",
             "00:00:10  Intro\n00:02:00  Interview proper"),
            (5, f"00:{split_at + 1:02d}:00", "00:20:00", "Street vox", "")]
    for r, s, e, text, segs in rows:
        ws.cell(r, 4, s); ws.cell(r, 5, e); ws.cell(r, 6, text); ws.cell(r, 7, segs)
    wb["Videos"].cell(4, 13, "no"); wb["Videos"].cell(4, 14, "tape is mislabelled")
    wb.save(path)


@pytest.fixture
def two_annotators(tmp_path):
    a, b = tmp_path / "NWFA_annotation_A.xlsx", tmp_path / "NWFA_annotation_B.xlsx"
    make_annotator(a, "Ann", 9); make_annotator(b, "Ben", 11)
    return [c.load_annotator(a), c.load_annotator(b)]


def test_annotator_reader(two_annotators):
    ann = two_annotators[0]
    assert ann["name"] == "Ann"
    progs = ann["programmes"]["N_6966411"]
    assert [(p["start"], p["end"]) for p in progs] == [(10, 540), (600, 1200)]
    assert progs[0]["segments"] == [(10, "Intro"), (120, "Interview proper")]
    assert ann["answers"]["N_6966411"]["label_matches"] == "no"


def test_two_annotators_give_four_extra_tracks(two_annotators):
    preds = c.load_predictions(c.META / "predictions-10.v3.xlsx")
    cat = c.load_catalogue(TEMPLATE)["N_6966411"]
    rec = c.build_record("N_6966411", cat, preds["N_6966411"], two_annotators, {}, None)
    names = [t["name"] for t in rec["metadata"]["media"]["tracks"]]
    assert names[:4] == ["human:Annotators · programmes (consensus)", "human:Annotators · segments (consensus)",
                         "pipe:Pipeline · programmes", "pipe:Pipeline · segments"]
    # breakdown tracks come last, so the UI can fold them away
    assert names[4:] == ["ann:Ann · programmes", "ann:Ann · segments", "ann:Ben · programmes", "ann:Ben · segments"]
    assert len(rec["ground_truth"]["annotator_notes"]) == 2


def test_consensus_marks_agreed_and_contested(two_annotators):
    # Ann: (10-540), (600-1200); Ben: (10-660), (720-1200): starts agree, first ends differ
    progs = c.cluster_programmes(two_annotators, "N_6966411")
    assert [(p["start"], p["end"], p["support"]) for p in progs] == [
        (10, 540, 1), (10, 660, 1), (600, 1200, 1), (720, 1200, 1)]
    assert {c.kind_of(p["support"], 2) for p in progs} == {"contested"}
    assert c.kind_of(1, 1) == ""


def test_identical_annotators_fully_agree(tmp_path):
    a, b = tmp_path / "NWFA_annotation_A.xlsx", tmp_path / "NWFA_annotation_B.xlsx"
    make_annotator(a, "Ann", 9); make_annotator(b, "Ben", 9)
    ann = [c.load_annotator(a), c.load_annotator(b)]
    progs = c.cluster_programmes(ann, "N_6966411")
    assert [p["support"] for p in progs] == [2, 2]
    assert {c.kind_of(p["support"], 2) for p in progs} == {"agreed"}


def test_pipeline_gaps_are_separators():
    progs = [{"start": 10, "end": 100}, {"start": 130, "end": 200}]
    seps = c.separators(progs, 200)
    assert [(s["start"], s["end"]) for s in seps] == [(0, 10), (100, 130)]
    assert all(s["text"].startswith("@sep ") for s in seps)


def test_judge_statuses():
    assert c.judge("year", "1915 1916 1917", "1916") == "match"
    assert c.judge("year", "1924", "1927") == "close"
    assert c.judge("colour", "b/w", "black-and-white") == "match"
    assert c.judge("colour", "b/w", "colour") == "differs"
    assert c.judge("year", None, "2007") == "added"
    assert c.judge("year", "1924", None) == "missing"
    assert c.judge("year", None, None) == "none"


def test_label_matches_join_on_reference_and_prog():
    m = c.load_label_matches(c.LABELS_DOCX)
    entry = m[("N_6975915", "00.00.41-1547")]
    assert entry["ranked"] and entry["true"]


def test_synopsis_is_neutral_text():
    assert c.judge("synopsis", "A mill in Bolton.", "Workers leave a Bolton mill.") == "text"
    assert c.judge("synopsis", None, "A mill.") == "added"
    assert c.judge("keywords", "mill; Bolton", "factory; Bolton") == "text"
    assert c.judge("category", "Sport", "Sport; Local") == "text"
